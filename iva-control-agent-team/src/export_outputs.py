"""Esportazione degli output del controllo: Excel, JSON e report Markdown.

``openpyxl`` e' una dipendenza opzionale: se manca, gli Excel sono sostituiti da
CSV equivalenti e il fatto viene dichiarato nel report, mai taciuto.
"""

from __future__ import annotations

import csv
import datetime as _dt
import json
import re
from pathlib import Path
from typing import Any, Sequence

from normalize import (
    Anomalia,
    ESITO_SOSPENDERE,
    GRAVITA_BLOCCANTE,
    GRAVITA_INFO,
    GRAVITA_VERIFICA,
    euro,
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - dipende dall'ambiente
    Workbook = None

MAX_ANOMALIE_REPORT = 15
_INTESTAZIONE_FILL = "1F3864"


def excel_disponibile() -> bool:
    return Workbook is not None


# ---------------------------------------------------------------------------
# Excel / CSV
# ---------------------------------------------------------------------------

def _scrivi_csv(percorso: Path, intestazioni: Sequence[str], righe: Sequence[Sequence[Any]]) -> Path:
    with percorso.open("w", encoding="utf-8-sig", newline="") as handle:
        scrittore = csv.writer(handle, delimiter=";")
        scrittore.writerow(intestazioni)
        scrittore.writerows(righe)
    return percorso


def scrivi_workbook(percorso: Path, fogli: dict[str, tuple[Sequence[str], Sequence[Sequence[Any]]]]) -> list[Path]:
    """Scrive un workbook multi-foglio; senza openpyxl produce un CSV per foglio."""
    if Workbook is None:
        prodotti = []
        for nome, (intestazioni, righe) in fogli.items():
            alternativo = percorso.with_name(f"{percorso.stem}_{nome}.csv")
            prodotti.append(_scrivi_csv(alternativo, intestazioni, righe))
        return prodotti

    workbook = Workbook()
    workbook.remove(workbook.active)
    for nome, (intestazioni, righe) in fogli.items():
        ws = workbook.create_sheet(nome[:31])
        ws.append(list(intestazioni))
        for riga in righe:
            ws.append(list(riga))
        fill = PatternFill("solid", fgColor=_INTESTAZIONE_FILL)
        for indice in range(1, len(intestazioni) + 1):
            cella = ws.cell(row=1, column=indice)
            cella.font = Font(bold=True, color="FFFFFF")
            cella.fill = fill
            cella.alignment = Alignment(vertical="center", wrap_text=True)
        for indice, titolo in enumerate(intestazioni, start=1):
            larghezza = max(12, min(55, len(str(titolo)) + 4))
            for riga in righe[:200]:
                if indice - 1 < len(riga):
                    larghezza = max(larghezza, min(55, len(str(riga[indice - 1])) + 2))
            ws.column_dimensions[get_column_letter(indice)].width = larghezza
        ws.freeze_panes = "A2"
        if righe:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(intestazioni))}1"
    workbook.save(percorso)
    return [percorso]


def esporta_anomalie(percorso: Path, anomalie: Sequence[Anomalia]) -> list[Path]:
    intestazioni = ["Codice", "Gravita", "Fonte", "Descrizione", "Impatto IVA",
                    "Azione proposta", "Stato", "Agente", "Note studio"]
    ordine = {GRAVITA_BLOCCANTE: 0, GRAVITA_VERIFICA: 1, GRAVITA_INFO: 2}
    ordinate = sorted(anomalie, key=lambda a: (ordine.get(a.gravita, 9), a.codice))
    righe = [[a.as_dict()[c] for c in intestazioni] for a in ordinate]
    return scrivi_workbook(percorso, {"Anomalie": (intestazioni, righe)})


CAMPI_DATASET = [
    "origine", "tipo_documento", "numero", "data", "data_registrazione", "data_ricezione",
    "controparte", "piva", "cf", "imponibile", "aliquota", "imposta", "natura",
    "esigibilita", "split_payment", "reverse_charge", "indetraibile", "totale_documento",
    "sezionale", "descrizione", "competenza", "fonte_file", "fonte_foglio", "fonte_riga",
    "fonte_pagina", "fonte_numero_documento", "note",
]


def esporta_dataset(percorso: Path, dataset: dict[str, list], extra: dict | None = None) -> list[Path]:
    """Esporta i dataset normalizzati, un foglio per famiglia di fonte."""
    fogli: dict[str, tuple[Sequence[str], Sequence[Sequence[Any]]]] = {}
    for nome, righe in dataset.items():
        if not righe:
            continue
        valori = [r.as_dict() for r in righe]
        fogli[nome] = (CAMPI_DATASET, [[v.get(c) for c in CAMPI_DATASET] for v in valori])
    for nome, (intestazioni, righe) in (extra or {}).items():
        fogli[nome] = (intestazioni, righe)
    if not fogli:
        fogli["Vuoto"] = (["nota"], [["Nessun dato normalizzato: verificare le fonti in ingresso."]])
    return scrivi_workbook(percorso, fogli)


def esporta_fonti(percorso: Path, audit: dict) -> Path:
    percorso.write_text(json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return percorso


# ---------------------------------------------------------------------------
# Report Markdown
# ---------------------------------------------------------------------------

def _riga_tabella(valori: Sequence[str]) -> str:
    return "| " + " | ".join(valori) + " |"


def _tabella_anomalie(anomalie: Sequence[Anomalia]) -> tuple[str, str]:
    rilevanti = [a for a in anomalie if a.gravita in (GRAVITA_BLOCCANTE, GRAVITA_VERIFICA)]
    ordine = {GRAVITA_BLOCCANTE: 0, GRAVITA_VERIFICA: 1}

    def peso_impatto(anomalia: Anomalia) -> float:
        testo = anomalia.impatto_iva.replace(".", "").replace(",", ".")
        try:
            return -float("".join(c for c in testo if c.isdigit() or c in ".-") or 0)
        except ValueError:
            return 0.0

    rilevanti.sort(key=lambda a: (ordine.get(a.gravita, 9), peso_impatto(a)))
    mostrate = rilevanti[:MAX_ANOMALIE_REPORT]
    if not mostrate:
        return "| — | — | — | Nessuna anomalia bloccante o da verificare rilevata. | — | — |", ""

    righe = [
        _riga_tabella([
            a.codice, a.gravita, a.fonte.replace("|", "/")[:120],
            a.descrizione.replace("|", "/").replace("\n", " "),
            a.impatto_iva, a.azione,
        ])
        for a in mostrate
    ]
    nota = ""
    if len(rilevanti) > MAX_ANOMALIE_REPORT:
        nota = (f"Sono riportate le prime {MAX_ANOMALIE_REPORT} anomalie per gravità e impatto; "
                f"le restanti {len(rilevanti) - MAX_ANOMALIE_REPORT} sono elencate in `anomalie_iva.xlsx`.")
    return "\n".join(righe), nota


def _elenco_informative(anomalie: Sequence[Anomalia]) -> str:
    informative = [a for a in anomalie if a.gravita == GRAVITA_INFO]
    if not informative:
        return "Nessun rilievo informativo."
    voci = [f"- `{a.codice}` {a.descrizione} *(fonte: {a.fonte})*" for a in informative[:20]]
    if len(informative) > 20:
        voci.append(f"- *(altri {len(informative) - 20} rilievi informativi in `anomalie_iva.xlsx`)*")
    return "\n".join(voci)


def componi_report(contesto: dict, template: Path | None = None) -> str:
    """Compone il report finale a partire dal contesto del controllo."""
    anomalie: list[Anomalia] = contesto["anomalie"]
    tabella_anomalie, nota_anomalie = _tabella_anomalie(anomalie)
    quadratura = contesto.get("quadratura", [])

    righe_quadratura = []
    for voce in quadratura:
        righe_quadratura.append(_riga_tabella([
            voce["voce"],
            euro(voce["da_fonti"]),
            euro(voce["da_bozza"]),
            euro(voce["differenza"]) if voce["differenza"] is not None else "n.d.",
        ]))

    sostituzioni = {
        "{{CLIENTE}}": contesto.get("cliente", "n.d."),
        "{{PERIODO}}": contesto.get("periodo_esteso", "n.d."),
        "{{ESITO}}": contesto.get("esito", ESITO_SOSPENDERE),
        "{{SINTESI}}": contesto.get("sintesi", ""),
        "{{FONTI}}": contesto.get("fonti_testo", ""),
        "{{FONTI_MANCANTI}}": contesto.get("fonti_mancanti_testo", "nessuna"),
        "{{TABELLA_ANOMALIE}}": tabella_anomalie,
        "{{NOTA_ANOMALIE}}": nota_anomalie,
        "{{RILIEVI_INFORMATIVI}}": _elenco_informative(anomalie),
        "{{RICHIESTE}}": contesto.get("richieste", "Nessuna richiesta al cliente."),
        "{{CONCLUSIONE}}": contesto.get("conclusione", ""),
        "{{NOTA_QUADRATURA}}": contesto.get("nota_quadratura", ""),
        "{{DATA_CONTROLLO}}": contesto.get("data_controllo", _dt.date.today().isoformat()),
    }

    if template and template.is_file():
        testo = template.read_text(encoding="utf-8")
        # La tabella di quadratura del template ha segnaposto per singola voce:
        # la sostituiamo in blocco con le righe effettivamente calcolate.
        inizio = testo.find("| Voce | Da fonti/registri")
        if inizio >= 0 and righe_quadratura:
            fine = testo.find("\n\n", inizio)
            intestazione = ("| Voce | Da fonti/registri | Da bozza liquidazione | Differenza |\n"
                            "|---|---:|---:|---:|\n")
            testo = testo[:inizio] + intestazione + "\n".join(righe_quadratura) + testo[fine:]
        testo = re.sub(r"<!--.*?-->\n?", "", testo, flags=re.DOTALL)
    else:
        testo = _report_predefinito()
        testo = testo.replace("{{TABELLA_QUADRATURA}}", "\n".join(righe_quadratura) or
                              "| — | — | — | — |")

    for segnaposto, valore in sostituzioni.items():
        testo = testo.replace(segnaposto, str(valore))
    return testo.strip() + "\n"


def _report_predefinito() -> str:
    return """# Controllo IVA — {{CLIENTE}} — {{PERIODO}}

## 1. Esito sintetico

Esito: **{{ESITO}}**

{{SINTESI}}

## 2. Fonti analizzate

{{FONTI}}

Fonti mancanti: {{FONTI_MANCANTI}}

## 3. Quadratura liquidazione

| Voce | Da fonti/registri | Da bozza liquidazione | Differenza |
|---|---:|---:|---:|
{{TABELLA_QUADRATURA}}

{{NOTA_QUADRATURA}}

## 4. Anomalie principali

| Codice | Gravità | Fonte | Descrizione | Impatto IVA | Azione |
|---|---|---|---|---:|---|
{{TABELLA_ANOMALIE}}

{{NOTA_ANOMALIE}}

## 5. Altri rilievi informativi

{{RILIEVI_INFORMATIVI}}

## 6. Documenti o chiarimenti da richiedere

{{RICHIESTE}}

## 7. Conclusione operativa

{{CONCLUSIONE}}

---

*Controllo eseguito il {{DATA_CONTROLLO}} con `iva-control-agent-team`. Dettaglio tecnico in
`anomalie_iva.xlsx` e `dataset_normalizzato.xlsx`; audit trail delle fonti in `fonti_utilizzate.json`.
Il presente report non sostituisce il giudizio professionale del commercialista.*
"""


def scrivi_report(percorso: Path, contesto: dict, template: Path | None = None) -> Path:
    percorso.write_text(componi_report(contesto, template), encoding="utf-8")
    return percorso


def scrivi_richiesta_documenti(percorso: Path, contesto: dict) -> Path | None:
    """Genera la richiesta al cliente, se ci sono documenti o chiarimenti da chiedere."""
    voci: list[str] = contesto.get("voci_richiesta", [])
    if not voci:
        return None
    testo = [
        f"# Richiesta documenti e chiarimenti — {contesto.get('cliente', 'n.d.')} — "
        f"{contesto.get('periodo_esteso', 'n.d.')}",
        "",
        "Per completare il controllo IVA del periodo sono necessari i seguenti documenti o chiarimenti.",
        "",
    ]
    for indice, voce in enumerate(voci, start=1):
        testo.append(f"{indice}. {voce}")
    testo += [
        "",
        "Fino alla ricezione di quanto sopra il controllo resta "
        f"**{contesto.get('esito', ESITO_SOSPENDERE)}**.",
        "",
        f"*Richiesta generata il {contesto.get('data_controllo', _dt.date.today().isoformat())} "
        "da `iva-control-agent-team`.*",
    ]
    percorso.write_text("\n".join(testo) + "\n", encoding="utf-8")
    return percorso
